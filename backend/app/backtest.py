"""Backtest engine + report assembly + Excel export."""
from __future__ import annotations

import datetime as dt
import io
import time

import numpy as np
import pandas as pd

from . import db
from .strategies.base import get_strategy


def _df_records(df: pd.DataFrame) -> list[dict]:
    out = df.to_dict(orient="records")
    for r in out:
        for k, v in list(r.items()):
            if isinstance(v, float) and pd.isna(v):
                r[k] = None
    return out


def load_bars(code: str, start: str, end: str) -> pd.DataFrame:
    """回测用的日线：`open/high/low/close` 是复权口径（价格 × 因子），另附 `raw_close`
    （交易所不复权收盘）与 `adj_factor`，量额保持交易所原值。

    信号必须用复权价，否则跨除权日会凭空亏掉一笔分红送转的市值；
    成交必须用不复权价，因为「一手 = 100 股 × 真实股价」才是真实占用资金，
    拿复权价算手数会把高价股撑成 6~7 倍、100 万连一手都买不起。
    """
    rows = db.query(
        "SELECT date, open*adj_factor AS open, high*adj_factor AS high, "
        "low*adj_factor AS low, close*adj_factor AS close, "
        "volume, amount, pct_chg, close AS raw_close, adj_factor "
        "FROM kline_daily WHERE code=? AND date>=? AND date<=? ORDER BY date",
        (code, start, end))
    return pd.DataFrame(rows)


def load_index_bars(index_code: str, start: str, end: str) -> list[dict]:
    rows = db.query(
        "SELECT index_code AS code, date, open, high, low, close, volume, amount, pct_chg "
        "FROM index_kline WHERE index_code=? AND date>=? AND date<=? ORDER BY date",
        (index_code, start, end))
    return rows


def synthesize_board_bars(board_code: str | None, industry: str | None,
                          start: str, end: str) -> list[dict]:
    """Industry index synthesized as equal-weighted cumulative return of
    locally available member stocks."""
    conds, args = [], []
    if board_code:
        conds.append("board_code=?")
        args.append(board_code)
    if industry:
        conds.append("industry=?")
        args.append(industry)
    if not conds:
        return []
    members = db.query(
        f"SELECT code FROM stocks WHERE {' OR '.join(conds)}", tuple(args))
    codes = sorted({m["code"] for m in members})
    if not codes:
        return []
    ph = ",".join("?" * len(codes))
    rows = db.query(
        f"SELECT k.code, k.date, k.close * k.adj_factor AS close FROM kline_daily k "
        f"WHERE k.code IN ({ph}) AND k.date>=? AND k.date<=? ORDER BY k.date",
        (*codes, start, end))
    if not rows:
        return []
    df = pd.DataFrame(rows).pivot(index="date", columns="code", values="close")
    df = df[df > 0].ffill(limit=5)
    df = df.dropna(axis=1, thresh=max(2, len(df) // 4))
    if df.empty or df.shape[1] == 0:
        return []
    # equal-weighted daily-return index of available members
    ret = df.pct_change(fill_method=None)
    level = 1000.0 * (1 + ret.mean(axis=1, skipna=True).fillna(0)).cumprod()
    out = []
    prev = None
    for date, val in level.items():
        if pd.isna(val):
            continue
        out.append({"date": date, "open": None, "high": None, "low": None,
                    "close": round(float(val), 2), "volume": None, "amount": None,
                    "pct_chg": round((float(val) / prev - 1) * 100, 4) if prev else None})
        prev = float(val)
    return out


def _range_pct(bars: list[dict]) -> float | None:
    closes = [b["close"] for b in bars if b.get("close")]
    if len(closes) < 2 or not closes[0]:
        return None
    return round((closes[-1] / closes[0] - 1) * 100, 2)


def walk_ledger(code: str, df: pd.DataFrame, signals: list[dict],
                capital: float, fee_rate: float) -> dict:
    """按K线重放一遍：信号在复权价上判定，成交与现金走不复权真实价。

    只在「信号日 + 复权因子变化日」推进状态机——其余交易日既不买卖、因子也没变，
    cash 与持股数都是常数，净值按段一次算完。逐根 iterrows 的写法全市场跑批要 20 分钟，
    这个写法不到 2 分钟，且与逐根版逐值等价（拿一批股票对账过）。
    单只报表和全市场跑批共用这一份，否则两边收益率对不上。
    """
    d = df.loc[df["adj_factor"].notna() & df["close"].notna() & df["raw_close"].notna()
               & (df["raw_close"] > 0), ["date", "close", "raw_close", "adj_factor"]]
    dates = d["date"].tolist()
    adj = d["close"].to_numpy(dtype=float)
    raw = d["raw_close"].to_numpy(dtype=float)
    fac = d["adj_factor"].to_numpy(dtype=float)
    n = len(dates)
    pos = {v: i for i, v in enumerate(dates)}

    sig: dict[int, str] = {}
    reason: dict[int, str] = {}
    for s in signals:
        i = pos.get(s["date"])
        if i is None:
            continue
        sig.setdefault(i, s["side"])
        reason[i] = s.get("reason", "")

    ex_div = set(int(i) for i in np.flatnonzero(fac[1:] != fac[:-1]) + 1)
    events = sorted(set(sig) | ex_div)

    cash = float(capital)
    qty = 0.0          # 真实持股数；除权除息日按因子折算后可能不是 100 的整数倍
    cost_basis = 0.0   # 本轮买入成交金额（真实元）
    buy_fee = 0.0
    trades: list[dict] = []
    markers: list[dict] = []
    missed = 0         # 因资金不够一手而被放弃的买入信号数，进 summary 别让它在静默里消失
    equity = np.empty(n)
    seg = 0
    for i in events:
        # 事件日之前的那段仍属于旧状态，先按旧状态结清，再改 cash/qty
        if i > seg:
            equity[seg:i] = cash + qty * raw[seg:i]
        # 除权除息：等比因子一变就按 F_t/F_(t-1) 折算持股数，等价「红利按除权日价再投」，
        # 于是 qty×不复权价 这条市值曲线在除权日不跳空，收益口径与复权价一致。
        if qty > 0 and fac[i] != fac[i - 1]:
            qty *= fac[i] / fac[i - 1]
        side = sig.get(i)
        if side == "buy" and qty == 0:
            lot = int(cash / (raw[i] * 100 * (1 + fee_rate))) * 100
            if lot >= 100:
                amount = lot * raw[i]
                fee = amount * fee_rate
                cash -= amount + fee
                qty, cost_basis, buy_fee = float(lot), amount, fee
                trades.append({
                    "date": dates[i], "side": "buy", "code": code, "qty": lot,
                    "price": round(raw[i], 3), "amount": round(amount, 2),
                    "fee": round(fee, 2), "pnl": None, "pnl_pct": None,
                    "reason": reason.get(i, ""),
                })
                # 买卖点画在复权K线上，所以标记价格用当日复权价，不能混真实价
                markers.append({"date": dates[i], "type": "buy", "price": round(adj[i], 3)})
            else:
                missed += 1
        elif side == "sell" and qty > 0:
            amount = qty * raw[i]
            fee = amount * fee_rate
            cash += amount - fee
            pnl = amount - fee - cost_basis - buy_fee
            trades.append({
                "date": dates[i], "side": "sell", "code": code, "qty": qty,
                "price": round(raw[i], 3), "amount": round(amount, 2), "fee": round(fee, 2),
                "pnl": round(pnl, 2),
                "pnl_pct": round(pnl / (cost_basis + buy_fee) * 100, 2)
                if cost_basis + buy_fee else None,
                "reason": reason.get(i, ""),
            })
            markers.append({"date": dates[i], "type": "sell", "price": round(adj[i], 3)})
            qty = cost_basis = buy_fee = 0.0
        equity[i] = cash + qty * raw[i]
        seg = i + 1
    if seg < n:
        equity[seg:n] = cash + qty * raw[seg:n]

    # 还持着的按最后一根不复权收盘强平
    final_equity = cash
    if qty > 0:
        amount = qty * raw[-1]
        fee = amount * fee_rate
        pnl = amount - fee - cost_basis - buy_fee
        trades.append({
            "date": dates[-1], "side": "sell", "code": code, "qty": qty,
            "price": round(raw[-1], 3), "amount": round(amount, 2), "fee": round(fee, 2),
            "pnl": round(pnl, 2),
            "pnl_pct": round(pnl / (cost_basis + buy_fee) * 100, 2)
            if cost_basis + buy_fee else None,
            "reason": "回测结束强制平仓",
        })
        markers.append({"date": dates[-1], "type": "sell", "price": round(adj[-1], 3)})
        final_equity = cash + amount - fee

    fees_total = sum(t["fee"] for t in trades)
    peak = np.maximum.accumulate(equity) if n else np.empty(0)
    return {
        "dates": dates, "equity": equity, "trades": trades, "markers": markers,
        "final_equity": final_equity, "pnl": final_equity - capital,
        "pnl_pct": (final_equity - capital) / capital * 100 if capital else None,
        "fees_total": fees_total,
        "order_count": len(trades),
        "trade_count": sum(1 for t in trades if t["side"] == "buy"),
        "missed": missed, "bars": n,
        "max_dd_pct": float(((equity - peak) / peak).min()) * 100 if n else None,
        "bh_pct": (adj[-1] / adj[0] - 1) * 100 if n else None,
    }


def run_backtest(code: str, start: str, end: str, capital: float,
                 fee_rate: float, strategy_id: str,
                 params: dict | None) -> dict:
    t0 = time.time()
    strategy = get_strategy(strategy_id)
    if strategy is None:
        raise ValueError(f"未知策略: {strategy_id}")
    df = load_bars(code, start, end)
    if df.empty or len(df) < 30:
        raise ValueError("该股票在区间内没有足够的本地K线数据，请先执行数据更新")
    # 缺复权因子（QMT 等比那份坏掉解不出）的这些根算不出复权价，静默丢掉会让持股段凭空断裂，
    # 所以直接拒绝并说清怎么补
    missing = df.loc[df["adj_factor"].isna(), "date"].tolist()
    if missing:
        raise ValueError(
            f"{code} 在 {start}~{end} 有 {len(missing)} 根缺复权因子（{missing[0]} 起），"
            "这些日子算不出复权价，回测结果不可信。请在 QMT 重导这只股票的「不复权」+"
            "「等比后复权」两份，再到「数据中心 → 本地数据（QMT 导入）」做一次增量导入即可自动补齐")
    led = walk_ledger(code, df, strategy.generate_signals(df, params or {}),
                      capital, fee_rate)
    trades, markers = led["trades"], led["markers"]
    missed = led["missed"]
    final_equity = led["final_equity"]
    fees_total = led["fees_total"]
    pnl_total = led["pnl"]

    stock = db.query_one("SELECT * FROM stocks WHERE code=?", (code,))
    from .datasource.eastmoney import index_secid_for, INDEXES
    idx_secid = index_secid_for(code)
    idx_code = next((k for k, v in INDEXES.items() if v[0] == idx_secid), "")
    idx_name = next((v[1] for v in INDEXES.values() if v[0] == idx_secid), idx_secid)
    index_bars = load_index_bars(idx_code, start, end)
    board_bars = synthesize_board_bars(
        stock["board_code"] if stock else None,
        stock["industry"] if stock else None, start, end)

    report_id = f"{code}_{strategy_id}_{dt.datetime.now().strftime('%Y%m%d%H%M%S')}"
    report = {
        "id": report_id,
        "code": code,
        "stock_name": stock["name"] if stock else code,
        "industry": stock["industry"] if stock else None,
        "board_code": stock["board_code"] if stock else None,
        "strategy": {"id": strategy_id, "name": strategy.name, "params": params or {}},
        "range": {"start": start, "end": end},
        "capital": capital,
        "fee_rate": fee_rate,
        "summary": {
            "elapsed_sec": round(time.time() - t0, 3),
            "order_count": len(trades),
            "trade_count": sum(1 for t in trades if t["side"] == "buy"),
            "missed_buy_signals": missed,
            "final_equity": round(final_equity, 2),
            "pnl": round(pnl_total, 2),
            "pnl_pct": round(pnl_total / capital * 100, 2) if capital else None,
            "fees_total": round(fees_total, 2),
            "index_name": idx_name,
            "index_pct": _range_pct(index_bars),
            "industry_pct": _range_pct(board_bars),
        },
        "stock_bars": _df_records(df),
        "markers": markers,
        "index_bars": index_bars,
        "board_bars": board_bars,
        "trades": trades,
    }
    _REPORTS[report_id] = report
    while len(_REPORTS) > 20:
        _REPORTS.pop(next(iter(_REPORTS)))
    return report


_REPORTS: dict[str, dict] = {}


def get_report(report_id: str) -> dict | None:
    return _REPORTS.get(report_id)


def export_excel(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    s = report["summary"]
    bold = Font(bold=True)
    rows = [
        ("股票代码", report["code"]), ("股票名称", report["stock_name"]),
        ("所属行业", report["industry"]),
        ("交易策略", report["strategy"]["name"]),
        ("回测区间", f"{report['range']['start']} ~ {report['range']['end']}"),
        ("初始资金", report["capital"]), ("手续费率", report["fee_rate"]),
        ("回测耗时(秒)", s["elapsed_sec"]), ("买入笔数", s["trade_count"]),
        ("总委托笔数", s["order_count"]),
        ("期末权益", s["final_equity"]),
        ("盈亏金额(含手续费)", s["pnl"]),
        ("盈亏比例(%)", s["pnl_pct"]), ("手续费合计", s["fees_total"]),
        (f"大盘涨跌幅({s['index_name']},%)", s["index_pct"]),
        ("行业涨跌幅(%)", s["industry_pct"]),
        ("信号因买不起一手被跳过", s["missed_buy_signals"]),
        ("成交口径", "不复权真实价（一手=100股）；除权除息日按复权因子折算持股数，等价红利再投"),
    ]
    for i, (k, v) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28

    ws2 = wb.create_sheet("交易明细")
    headers = ["交易日期", "交易股票", "方向", "交易数量(股)", "成交价(元)",
               "交易金额", "手续费", "盈亏金额", "盈亏比例(%)", "触发原因"]
    ws2.append(headers)
    for c in ws2[1]:
        c.font = bold
    code_name = f"{report['code']} {report['stock_name']}"
    for t in report["trades"]:
        ws2.append([t["date"], code_name, "买入" if t["side"] == "buy" else "卖出",
                    round(t["qty"], 2), t["price"], t["amount"], t["fee"],
                    t["pnl"], t["pnl_pct"], t["reason"]])
    widths = [12, 18, 8, 14, 12, 14, 10, 12, 12, 24]
    for i, w in enumerate(widths):
        ws2.column_dimensions[chr(ord("A") + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
